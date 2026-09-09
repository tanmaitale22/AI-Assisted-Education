from pptx import Presentation


def extract_ppt_text(file_path):
    presentation = Presentation(file_path)

    text = ""

    for slide in presentation.slides:
        for shape in slide.shapes:

            if hasattr(shape, "text"):
                text += shape.text + "\n"

        text += "\n"

    return text